import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# Title
title = doc.add_heading('Navigating the New UAE Visa Landscape in 2026: What You Need to Know', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Meta information
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('By Arihant Travel Team | January 20, 2026 | 18 min read').italic = True

doc.add_paragraph()

# Introduction
intro = doc.add_paragraph()
intro.add_run('The UAE has officially moved into a new era of "purpose-based" immigration. Gone are the days of a one-size-fits-all entry permit. In 2026, the system is designed to attract specific talents—from AI specialists to environmental activists—while making it easier than ever for tourists to explore the entire Gulf region.')
doc.add_paragraph('Whether you\'re planning a week-long holiday or a decade-long career, here is your definitive guide to the UAE\'s latest visa rules.')

# Section 1: GCC Unified Tourist Visa
doc.add_heading('1. The Big Game Changer: The GCC Unified Tourist Visa', 1)
doc.add_paragraph('If you\'ve ever dreamed of a road trip from Dubai to Muscat or Doha, 2026 is your year. The GCC Unified Tourist Visa (often called the "Schengen of the Middle East") is now fully operational.')

doc.add_heading('How it works:', 2)
doc.add_paragraph('A single visa allows you to travel across the UAE, Saudi Arabia, Qatar, Bahrain, Oman, and Kuwait.')

doc.add_heading('Who it\'s for:', 2)
doc.add_paragraph('Tourists and business travelers looking for multi-country itineraries without the hassle of multiple applications.')

doc.add_heading('Cost:', 2)
doc.add_paragraph('Projected to cost between USD $90 – $130 (AED 330 – 480) for multi-country access.')

# Section 2: Updated Visit Visas
doc.add_heading('2. Updated Visit Visas: More Time, Less Stress', 1)
doc.add_paragraph('The standard visit visa has received a major refresh to cater to modern travelers:')

updates = doc.add_paragraph()
updates.add_run('• Flexible Durations: ').bold = True
updates.add_run('You can now choose between 30, 60, or 90-day stays.\n')
updates.add_run('• In-Country Extensions: ').bold = True
updates.add_run('You no longer need to do "visa runs" to the border. Most visit visas can now be extended online via the ICP or GDRFA portals while you are still in the UAE.\n')
updates.add_run('• No More Grace Period: ').bold = True
updates.add_run('Be careful! The 10-day grace period has been removed. Overstaying now results in an immediate fine of AED 50 per day.')

# Section 3: New Niche Visa Categories
doc.add_heading('3. New "Niche" Visa Categories', 1)
doc.add_paragraph('The UAE has introduced several new entry permits to match emerging global industries:')

niche = doc.add_paragraph()
niche.add_run('• AI Specialist Visa: ').bold = True
niche.add_run('For engineers and researchers in artificial intelligence.\n')
niche.add_run('• Blue Visa: ').bold = True
niche.add_run('A 10-year residency for environmental contributors and sustainability researchers.\n')
niche.add_run('• Event Visa: ').bold = True
niche.add_run('Specifically for those attending international conferences, festivals, or sports events.\n')
niche.add_run('• Jobseeker Visa: ').bold = True
niche.add_run('No sponsor required for skilled professionals or graduates from top-tier universities to come and look for work.')

# Section 4: Golden Visa
doc.add_heading('4. The Golden Visa: 2026 Expansion', 1)
doc.add_paragraph('The Golden Visa remains the "holy grail" of residency, offering 10 years of stability. New categories for 2026 include:')

golden = doc.add_paragraph()
golden.add_run('• Nurses & Teachers: ').bold = True
golden.add_run('Recognizing long-term service (typically 15+ years) in the UAE.\n')
golden.add_run('• Humanitarian Contributors: ').bold = True
golden.add_run('For those who have made significant charitable impacts or donations.\n')
golden.add_run('• Content Creators: ').bold = True
golden.add_run('A structured pathway for influencers and digital artists through "Creators HQ."')

doc.add_heading('Cost Breakdown:', 2)
cost = doc.add_paragraph()
cost.add_run('• 10-Year Golden Visa (Property Investor): ').bold = True
cost.add_run('Total cost approximately AED 9,500 – 10,500\n')
cost.add_run('• 10-Year Golden Visa (Professionals/Managers/Scientists): ').bold = True
cost.add_run('Total cost approximately AED 4,000 – 4,800')

# Section 5: Quick Checklist
doc.add_heading('5. Quick Checklist for Applicants', 1)
doc.add_paragraph('Before you hit "submit" on your application, ensure you have these 2026 requirements ready:')

checklist = doc.add_paragraph()
checklist.add_run('• Passport Cover Page: ').bold = True
checklist.add_run('A scan of the external cover of your passport is now mandatory\n')
checklist.add_run('• Health Insurance: ').bold = True
checklist.add_run('Must be valid for the entire duration of your stay\n')
checklist.add_run('• Financial Proof: ').bold = True
checklist.add_run('Minimum bank balance around USD $4,000 for 5-year multi-entry visa')

# Add page break
doc.add_page_break()

# UAE Visit Visa Fee Guide 2026
doc.add_heading('UAE Visit Visa Fee Guide 2026', 1)
doc.add_paragraph('This table provides a comprehensive look at the estimated visa costs for 2026. Please note that fees can vary slightly depending on whether you apply through the official ICP/GDRFA portals, an airline (like Emirates or Etihad), or a travel agency.')

# Create table
table = doc.add_table(rows=10, cols=5)
table.style = 'Light Grid Accent 1'

# Header row
header_cells = table.rows[0].cells
header_cells[0].text = 'Visa Type'
header_cells[1].text = 'Entry Type'
header_cells[2].text = 'Duration'
header_cells[3].text = 'Fee (AED)'
header_cells[4].text = 'Fee (USD)'

# Make header bold
for cell in header_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True

# Data rows
data = [
    ['Transit Visa', 'Single', '48 Hours', 'Free / AED 40*', '~$10'],
    ['Transit Visa', 'Single', '96 Hours', 'AED 180', '~$49'],
    ['Tourist Visa', 'Single', '30 Days', 'AED 300 - 450', '~$80 - 120'],
    ['Tourist Visa', 'Multiple', '30 Days', 'AED 650 - 800', '~$175 - 220'],
    ['Tourist Visa', 'Single', '60 Days', 'AED 500 - 650', '~$135 - 175'],
    ['Tourist Visa', 'Multiple', '60 Days', 'AED 900 - 1,100', '~$245 - 300'],
    ['Tourist Visa', 'Single', '90 Days', 'AED 700 - 900', '~$190 - 245'],
    ['Tourist Visa', 'Multiple', '90 Days', 'AED 1,600 - 1,900', '~$435 - 520'],
    ['Job Exploration', 'Single', '60 Days', 'AED 500 - 600', '~$135 - 165']
]

for i, row_data in enumerate(data, start=1):
    cells = table.rows[i].cells
    for j, value in enumerate(row_data):
        cells[j].text = value

doc.add_paragraph()
note = doc.add_paragraph()
note.add_run('Note: ').bold = True
note.add_run('"Visa Change" or "Inside Country" status updates typically incur an additional fee of approximately AED 600 - 675 if you do not wish to exit and re-enter the country.')

# Long-Term & Residency Fees
doc.add_heading('Long-Term & Residency Fees (2026)', 1)
doc.add_paragraph('For those looking at long-term stays, the costs include the visa fee, medical fitness test, and the 10-year Emirates ID.')

longterm = doc.add_paragraph()
longterm.add_run('• 10-Year Golden Visa (Property Investor): ').bold = True
longterm.add_run('Total cost approximately AED 9,500 – 10,500. This includes the DLD (Dubai Land Department) fees and administrative costs.\n')
longterm.add_run('• 10-Year Golden Visa (Professionals/Managers/Scientists): ').bold = True
longterm.add_run('Total cost approximately AED 4,000 – 4,800.\n')
longterm.add_run('• 5-Year Green Visa (Freelancers/Skilled Employees): ').bold = True
longterm.add_run('Total cost approximately AED 2,500 – 3,500.\n')
longterm.add_run('• GCC Unified Tourist Visa (New for 2026): ').bold = True
longterm.add_run('Projected to cost between USD $90 – $130 (AED 330 – 480) for multi-country access.')

# Mandatory Overstay & Insurance Rules
doc.add_heading('Mandatory Overstay & Insurance Rules', 1)

overstay = doc.add_paragraph()
overstay.add_run('• Standardized Fines: ').bold = True
overstay.add_run('The grace period has been abolished. Overstaying any visa now costs exactly AED 50 per day.\n')
overstay.add_run('• Health Insurance: ').bold = True
overstay.add_run('You cannot process a 2026 visa without valid travel insurance. Basic policies for a 30-day stay start as low as AED 40 - 60.')

# Add page break
doc.add_page_break()

# Indian Citizens Visa Requirements
doc.add_heading('Indian Citizens Visa Requirements', 1)
doc.add_paragraph('For Indian citizens, the UAE has become one of the most accessible destinations in 2026, with several streamlined visa options ranging from short-term stopovers to long-term residency.')

doc.add_heading('1. Visa-on-Arrival (VoA) for Eligible Indians', 2)
doc.add_paragraph('This is the fastest way to enter. If you hold specific international visas, you can get your entry permit directly at the airport.')

doc.add_heading('Eligibility Requirements:', 3)
doc.add_paragraph('You must have a regular Indian passport valid for at least 6 months AND one of the following (valid for at least 6 months):')

eligibility = doc.add_paragraph()
eligibility.add_run('• A US Visit Visa or Green Card\n')
eligibility.add_run('• A UK Residence Permit\n')
eligibility.add_run('• An EU/Schengen Residence Permit\n\n')
eligibility.add_run('Duration: ').bold = True
eligibility.add_run('14 days (Single Entry)\n')
eligibility.add_run('Extension: ').bold = True
eligibility.add_run('Can be extended once for an additional 14 days\n')
eligibility.add_run('Estimated Fee: ').bold = True
eligibility.add_run('Approximately AED 100 - 120 ($27–33)')

doc.add_heading('2. Standard Pre-Arranged Tourist Visas', 2)
doc.add_paragraph('If you do not meet the VoA criteria, you must apply in advance through a travel agency, an airline (Emirates, Etihad, FlyDubai, Air Arabia), or a hotel.')

# Create table for Indian visa types
indian_table = doc.add_table(rows=6, cols=4)
indian_table.style = 'Light Grid Accent 1'

# Header
header = indian_table.rows[0].cells
header[0].text = 'Visa Type'
header[1].text = 'Validity (Stay)'
header[2].text = 'Best For...'
header[3].text = 'Estimated Cost (AED)'

for cell in header:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True

# Data
indian_data = [
    ['48-Hour Transit', '48 Hours', 'Short layovers', '~AED 40 - 50'],
    ['96-Hour Transit', '96 Hours', 'Quick city stopovers', '~AED 180 - 200'],
    ['30-Day Single Entry', '30 Days', 'Standard holidays', '~AED 300 - 450'],
    ['60-Day Single Entry', '60 Days', 'Family visits', '~AED 500 - 650'],
    ['30/60-Day Multiple', '30 or 60 Days', 'Regional GCC travel', '~AED 650 - 1,100']
]

for i, row_data in enumerate(indian_data, start=1):
    cells = indian_table.rows[i].cells
    for j, value in enumerate(row_data):
        cells[j].text = value

doc.add_heading('3. The 5-Year Multiple Entry Tourist Visa', 2)
doc.add_paragraph('This is a "self-sponsored" visa specifically designed for frequent travelers. It does not require a UAE-based guarantor.')

doc.add_heading('Key Rules:', 3)
five_year = doc.add_paragraph()
five_year.add_run('• Stay up to 90 days per visit, extendable for another 90 days (max 180 days per year)\n')
five_year.add_run('• Mandatory Financial Requirement: ').bold = True
five_year.add_run('You must provide a 6-month bank statement showing a minimum balance of USD $4,000 (roughly ₹3.3 Lakhs)\n')
five_year.add_run('• Estimated Total Cost: ').bold = True
five_year.add_run('Approximately AED 650 – 800 (excluding the refundable security deposit of ~AED 3,000 which some portals may require)')

# Add page break
doc.add_page_break()

# Required Documents Checklist
doc.add_heading('Required Documents Checklist', 1)
doc.add_paragraph('To ensure your application is processed within the typical 3–5 working days, have these ready in high-quality digital scans:')

doc.add_heading('Passport Documents:', 2)
passport_docs = doc.add_paragraph()
passport_docs.add_run('• Clear color copies of the first page (bio-data) and the last page (address)\n')
passport_docs.add_run('• A scan of the outside front cover of your passport (newly mandatory in 2026)\n')
passport_docs.add_run('• Passport must be valid for at least 6 months from date of entry')

doc.add_heading('Photograph:', 2)
photo = doc.add_paragraph()
photo.add_run('• Recent passport-sized color photo with white background\n')
photo.add_run('• Face must occupy 70–80% of frame\n')
photo.add_run('• Photo must be taken within last 6 months')

doc.add_heading('Travel Proof:', 2)
travel = doc.add_paragraph()
travel.add_run('• Confirmed return flight ticket\n')
travel.add_run('• Hotel voucher or host\'s Emirates ID and invitation letter')

doc.add_heading('Insurance & Financial:', 2)
insurance = doc.add_paragraph()
insurance.add_run('• Travel insurance with minimum USD $30,000 emergency medical coverage\n')
insurance.add_run('• Bank statement showing balance of at least AED 5,000 (for 60/90-day visas)')

doc.add_heading('Important Notes for 2026:', 2)
notes = doc.add_paragraph()
notes.add_run('• No Grace Period: ').bold = True
notes.add_run('Ensure you exit or renew your visa before the expiry date. Fines of AED 50 per day apply starting the very first day of overstay.\n')
notes.add_run('• Processing Time: ').bold = True
notes.add_run('Standard processing takes 3–5 working days. Express service is available (24–48 hours) for an additional fee.\n')
notes.add_run('• Scan Quality: ').bold = True
notes.add_run('Do not use mobile "photo-of-a-photo" scans. Only use high-quality digital scans to avoid automatic rejection by the AI-driven ICP system.')

# Add page break
doc.add_page_break()

# FAQ Section
doc.add_heading('Frequently Asked Questions: UAE Visa 2026', 1)

doc.add_heading('Q: Is there still a 10-day grace period for tourist visas in 2026?', 2)
doc.add_paragraph('A: No. The UAE has officially removed the 10-day grace period. Fines of AED 50 per day now begin the very first day after your visa expires. It is highly recommended to apply for an extension or exit the country at least 2–3 days before your expiry date.')

doc.add_heading('Q: Can I extend my tourist visa without leaving the UAE?', 2)
doc.add_paragraph('A: Yes! One of the best updates for 2026 is the ability to extend 30, 60, or 90-day visas entirely online via the ICP or GDRFA portals. You no longer need to do "visa runs" to the border.')

doc.add_heading('Q: What is the "GCC Unified Tourist Visa" I keep hearing about?', 2)
doc.add_paragraph('A: Often called the "Grand Tours Visa," this is a new single-visa system launched in 2026 that allows travelers to visit all six GCC countries (UAE, Saudi Arabia, Qatar, Kuwait, Oman, and Bahrain) using one permit, similar to a Schengen visa.')

doc.add_heading('Q: How much bank balance do I need for a 5-year multiple-entry visa?', 2)
doc.add_paragraph('A: For the self-sponsored 5-year tourist visa, applicants generally need to provide a bank statement for the last 6 months showing a minimum balance of USD $4,000 (approx. AED 14,700) or its equivalent in other currencies.')

doc.add_heading('Q: Can I sponsor my family if I am on a visitor visa?', 2)
sponsor = doc.add_paragraph('A: No, you must have a valid Residency Visa to sponsor family members. In 2026, the salary requirements for residents sponsoring visitors have been updated:')
sponsor_list = doc.add_paragraph()
sponsor_list.add_run('• AED 4,000/month for first-degree relatives (spouse/children)\n')
sponsor_list.add_run('• AED 8,000/month for second-degree relatives (siblings/grandparents)\n')
sponsor_list.add_run('• AED 15,000/month to sponsor a friend')

doc.add_heading('Q: What are the mandatory documents for a UAE visa in 2026?', 2)
mandatory = doc.add_paragraph('A: In addition to your passport and photo, the 2026 system now strictly requires:')
mandatory_list = doc.add_paragraph()
mandatory_list.add_run('1. A copy of the external passport cover\n')
mandatory_list.add_run('2. Confirmed accommodation proof (hotel booking or host\'s address)\n')
mandatory_list.add_run('3. Valid health/travel insurance for the duration of your stay')

# Final Thoughts
doc.add_heading('Final Thoughts', 1)
doc.add_paragraph('The UAE visa landscape in 2026 offers unprecedented flexibility and opportunities for travelers, professionals, and investors alike. From the revolutionary GCC Unified Tourist Visa to the expanded Golden Visa categories, the Emirates continues to position itself as a global hub for tourism, business, and innovation.')
doc.add_paragraph('Whether you\'re planning a quick weekend getaway or considering long-term residency, understanding these new regulations will help you navigate the process smoothly and avoid costly mistakes.')

# Contact Information
doc.add_heading('Need Help with Your UAE Visa Application?', 1)
contact = doc.add_paragraph('Let Arihant Travel handle your UAE visa application with expert guidance and hassle-free processing. We specialize in all visa types including tourist, transit, and long-term residency visas.')
contact.add_run('\n\nVisit: https://arihantlink.com/uae-visa').bold = True
contact.add_run('\nWhatsApp: +971 58 594 5007').bold = True

# Hashtags
doc.add_heading('Social Media Hashtags:', 1)
hashtags = doc.add_paragraph('#UAEVisa2026 #DubaiTravel #GCCGrandTours #VisitDubai #UAEGoldenVisa #TravelTips2026 #MyDubai #DubaiLife #UAEImmigration #ArihantTravels')

# Save the document
output_path = '/Users/neerajjain/Documents/Arihantlink/latest-web/blog/UAE_Visa_2026_Guide.docx'
doc.save(output_path)
print(f"Document saved successfully to: {output_path}")
